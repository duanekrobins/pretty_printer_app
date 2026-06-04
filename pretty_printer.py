#!/usr/bin/env python3
"""
TestSavvy XML / Excel / JSON Pretty Printer and Markdown Documentation Generator

Developer: Duane K Robinson
Organization: State of Utah
Date: March 11th 2026

Program Purpose
---------------
This Python program converts machine-generated automation artifacts into
human-readable Markdown documentation suitable for review, troubleshooting,
knowledge-base storage, and AI-assisted ingestion tools such as NotebookLM.

The program is configuration-driven. Runtime behavior is controlled by a JSON
configuration file passed with the --config command-line argument.

Primary Capabilities
--------------------
1. XML Processing
   - Reads XML files from configured input directories.
   - Supports raw XML pretty-print output.
   - Supports enhanced TestSavvy execution report output.
   - Extracts TestSavvy automation_sequence metadata.
   - Maps automation_sequence id to both Test Case ID and Automation Sequence ID.
   - Extracts scenario metadata, scenario attributes, step metadata, runtime
     variables, machine settings, execution summary fields, and raw attributes.
   - Optionally appends a full XML Inventory section listing every parsed node,
     node path, depth, attributes, and text.

2. Markdown Generation
   - Writes Markdown files using the configured output extension.
   - Supports one output Markdown file per input file.
   - Supports combined Markdown output for multiple input files.
   - Structures TestSavvy XML data into sections suitable for NotebookLM,
     repository documentation, and human QA review.

3. Excel Processing
   - Reads .xlsx and .xlsm workbooks.
   - Converts each worksheet/tab into a Markdown table.
   - Supports one Markdown output per worksheet.
   - Supports combined output when configured.

4. JSON Processing
   - Pretty prints JSON.
   - Can wrap pretty-printed JSON in Markdown code fences when output is .md.

Recommended NotebookLM Use
--------------------------
For NotebookLM ingestion, use:
- xml_output_mode: "testsavvy_execution_report"
- include_xml_inventory: true
- output_extension: ".md"

This gives NotebookLM both a human-readable execution report and a complete
node/attribute inventory so XML details are not silently omitted.

Typical Command
---------------
python pretty_printer.py --config pretty_print_config.json

Important Notes
---------------
- This utility is intentionally read-only with respect to source files.
- Output files are written only to the configured output directory.
- overwrite=true allows regenerated output to replace previous output.
- The parser is optimized for TestSavvy execution XML payloads but also retains
  raw XML output support for debugging and exact source review.
"""

from __future__ import annotations

import argparse
import fnmatch
import json
import re
import io
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable, List, Optional, Tuple, Dict, Any
import xml.etree.ElementTree as ET

from openpyxl import load_workbook


@dataclass(frozen=True)
class AppConfig:
    input_dirs: List[Path]
    output_dir: Path
    file_types: Tuple[str, ...]
    recursive: bool
    overwrite: bool
    indent: int
    json_sort_keys: bool
    preserve_xml_declaration: bool
    xml_encoding: str
    output_structure: str
    output_extension: Optional[str]
    exclude_patterns: Tuple[str, ...]
    xml_markdown_mode: str
    combined_output_filename: str
    xml_output_mode: str
    excel_markdown_mode: str
    excel_include_empty_sheets: bool
    excel_max_cell_length: int
    excel_sheet_name_in_filename: bool
    include_xml_inventory: bool


def _normalize_extension(ext: Optional[str]) -> Optional[str]:
    """
    Normalize a configured output extension so values like 'md' and '.md' behave consistently.
    """
    if ext is None:
        return None
    ext = str(ext).strip()
    if ext == "":
        return ""
    if not ext.startswith("."):
        ext = "." + ext
    return ext.lower()



# -----------------------------------------------------------------------------
# Configuration loading and validation
# -----------------------------------------------------------------------------
def load_config(config_path: Path) -> AppConfig:
    """
    Read the JSON configuration file, validate key options, normalize paths and extensions, and return an immutable AppConfig object.
    """
    if not config_path.exists():
        raise FileNotFoundError(f"Config file not found: {config_path}")

    raw = json.loads(config_path.read_text(encoding="utf-8"))

    input_dirs = [Path(p).expanduser().resolve() for p in raw.get("input_dirs", [])]
    if not input_dirs:
        raise ValueError("Config must include 'input_dirs' with at least one directory.")

    output_dir = Path(raw.get("output_dir", "./output")).expanduser().resolve()

    file_types_raw = raw.get("file_types", ["xml", "xlsx", "xlsm"])
    file_types = tuple(ft.lower().strip(".") for ft in file_types_raw)
    allowed = {"json", "xml", "xlsx", "xlsm"}
    unknown = [ft for ft in file_types if ft not in allowed]
    if unknown:
        raise ValueError(f"Unknown file types in config: {unknown}. Allowed: {sorted(allowed)}")

    recursive = bool(raw.get("recursive", True))
    overwrite = bool(raw.get("overwrite", False))
    indent = int(raw.get("indent", 4))
    json_sort_keys = bool(raw.get("json_sort_keys", True))
    preserve_xml_declaration = bool(raw.get("preserve_xml_declaration", True))
    xml_encoding = str(raw.get("xml_encoding", "utf-8"))

    output_structure = str(raw.get("output_structure", "mirror")).lower()
    if output_structure not in {"mirror", "flat"}:
        raise ValueError("output_structure must be 'mirror' or 'flat'.")

    output_extension = _normalize_extension(raw.get("output_extension", ".md"))
    exclude_patterns = tuple(raw.get("exclude_patterns", []))

    xml_markdown_mode = str(raw.get("xml_markdown_mode", "per_file")).lower()
    if xml_markdown_mode not in {"per_file", "combined"}:
        raise ValueError("xml_markdown_mode must be 'per_file' or 'combined'.")

    xml_output_mode = str(raw.get("xml_output_mode", "testsavvy_execution_report")).lower()
    if xml_output_mode not in {"testsavvy_execution_report", "raw_markdown"}:
        raise ValueError("xml_output_mode must be 'testsavvy_execution_report' or 'raw_markdown'.")

    combined_output_filename = str(raw.get("combined_output_filename", "combined_output.md"))

    excel_markdown_mode = str(raw.get("excel_markdown_mode", "per_sheet")).lower()
    if excel_markdown_mode not in {"per_sheet"}:
        raise ValueError("excel_markdown_mode currently supports only 'per_sheet'.")

    include_xml_inventory = bool(raw.get("include_xml_inventory", True))

    return AppConfig(
        input_dirs=input_dirs,
        output_dir=output_dir,
        file_types=file_types,
        recursive=recursive,
        overwrite=overwrite,
        indent=indent,
        json_sort_keys=json_sort_keys,
        preserve_xml_declaration=preserve_xml_declaration,
        xml_encoding=xml_encoding,
        output_structure=output_structure,
        output_extension=output_extension,
        exclude_patterns=exclude_patterns,
        xml_markdown_mode=xml_markdown_mode,
        combined_output_filename=combined_output_filename,
        xml_output_mode=xml_output_mode,
        excel_markdown_mode=excel_markdown_mode,
        excel_include_empty_sheets=bool(raw.get("excel_include_empty_sheets", False)),
        excel_max_cell_length=int(raw.get("excel_max_cell_length", 500)),
        excel_sheet_name_in_filename=bool(raw.get("excel_sheet_name_in_filename", True)),
        include_xml_inventory=include_xml_inventory,
    )


def is_excluded(path: Path, patterns: Tuple[str, ...]) -> bool:
    """
    Return True when a source file name matches one of the configured exclude glob patterns.
    """
    return any(fnmatch.fnmatch(path.name, pat) for pat in patterns) if patterns else False


def iter_files(input_dir: Path, recursive: bool, file_types: Tuple[str, ...]) -> Iterable[Path]:
    """
    Yield source files from an input directory according to recursion and configured file extensions.
    """
    if not input_dir.exists():
        return
    walker = input_dir.rglob if recursive else input_dir.glob
    seen = set()
    for ft in file_types:
        for path in walker(f"*.{ft}"):
            if path not in seen:
                seen.add(path)
                yield path


def apply_output_extension(filename: str, forced_ext: Optional[str]) -> str:
    """
    Apply the configured output extension to a file name while preserving the base stem.
    """
    if forced_ext is None or forced_ext == "":
        return filename
    return Path(filename).with_suffix(forced_ext).name


def compute_output_path(src: Path, input_root: Path, output_dir: Path, output_structure: str, forced_ext: Optional[str]) -> Path:
    """
    Build the destination output path for an input file using either mirror or flat output structure.
    """
    if output_structure == "flat":
        return output_dir / apply_output_extension(src.name, forced_ext)
    rel = src.relative_to(input_root)
    return output_dir / rel.with_name(apply_output_extension(rel.name, forced_ext))


def ensure_parent_dir(path: Path) -> None:
    """
    Create the parent directory for an output file if it does not already exist.
    """
    path.parent.mkdir(parents=True, exist_ok=True)


def pretty_print_json(src: Path, indent: int, sort_keys: bool) -> str:
    """
    Load JSON and serialize it back out with configured indentation and optional sorted keys.
    """
    text = src.read_text(encoding="utf-8")
    obj = json.loads(text)
    return json.dumps(obj, indent=indent, ensure_ascii=False, sort_keys=sort_keys) + "\n"


def pretty_print_xml(src: Path, indent: int, encoding: str, preserve_declaration: bool) -> bytes:
    """
    Parse XML, apply indentation, and serialize it back to bytes using the configured XML encoding and declaration preference.
    """
    tree = ET.parse(str(src), parser=ET.XMLParser())
    ET.indent(tree, space=" " * indent)
    buf = io.BytesIO()
    tree.write(
        buf,
        encoding=encoding,
        xml_declaration=preserve_declaration,
        short_empty_elements=True,
    )
    buf.write(b"\n")
    return buf.getvalue()


def wrap_xml_markdown(src: Path, xml_text: str) -> str:
    """
    Wrap pretty-printed XML text inside a Markdown XML code fence.
    """
    return f"# Pretty Printed XML\n\n**Source file:** `{src.name}`\n\n```xml\n{xml_text.rstrip()}\n```\n"


def wrap_json_markdown(src: Path, json_text: str) -> str:
    """
    Wrap pretty-printed JSON text inside a Markdown JSON code fence.
    """
    return f"# Pretty Printed JSON\n\n**Source file:** `{src.name}`\n\n```json\n{json_text.rstrip()}\n```\n"


def sanitize_sheet_name(sheet_name: str) -> str:
    """
    Convert an Excel worksheet name into a safe filename component.
    """
    value = re.sub(r"[^\w\-. ]+", "_", sheet_name.strip())
    value = re.sub(r"\s+", "_", value)
    return value or "Sheet"


def markdown_escape(value: object, max_len: int = 500) -> str:
    """
    Escape Markdown table-sensitive characters and normalize whitespace so cell content renders safely.
    """
    text = "" if value is None else str(value)
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\n", "<br>")
    text = text.replace("|", "\\|")
    text = re.sub(r"\s+", " ", text).strip()
    if max_len > 0 and len(text) > max_len:
        text = text[: max_len - 3] + "..."
    return text


def worksheet_has_data(ws) -> bool:
    """
    Check whether an Excel worksheet contains at least one non-empty cell.
    """
    for row in ws.iter_rows(values_only=True):
        if any(cell not in (None, "") for cell in row):
            return True
    return False



# -----------------------------------------------------------------------------
# Excel worksheet to Markdown conversion
# -----------------------------------------------------------------------------
def worksheet_to_markdown(src: Path, ws, max_cell_length: int) -> str:
    """
    Convert an Excel worksheet into a Markdown table using the first non-empty row as the header row.
    """
    rows = list(ws.iter_rows(values_only=True))
    non_empty_rows = []
    max_cols = 0

    for row in rows:
        trimmed = list(row)
        while trimmed and trimmed[-1] in (None, ""):
            trimmed.pop()
        if trimmed:
            non_empty_rows.append(trimmed)
            max_cols = max(max_cols, len(trimmed))

    if not non_empty_rows:
        return (
            f"# Excel Worksheet Export\n\n"
            f"**Workbook:** `{src.name}`\n\n"
            f"**Worksheet:** `{ws.title}`\n\n"
            f"_This worksheet is empty._\n"
        )

    normalized = []
    for row in non_empty_rows:
        padded = list(row) + [""] * (max_cols - len(row))
        normalized.append([markdown_escape(cell, max_cell_length) for cell in padded])

    header = normalized[0]
    if all(h == "" for h in header):
        header = [f"Column {i+1}" for i in range(max_cols)]
        data_rows = normalized
    else:
        header = [h if h else f"Column {i+1}" for i, h in enumerate(header)]
        data_rows = normalized[1:]

    lines = [
        "# Excel Worksheet Export",
        "",
        f"**Workbook:** `{src.name}`",
        "",
        f"**Worksheet:** `{ws.title}`",
        "",
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * len(header)) + " |",
    ]

    for row in data_rows:
        lines.append("| " + " | ".join(row) + " |")

    if not data_rows:
        lines.append("| " + " | ".join([""] * len(header)) + " |")

    lines.append("")
    return "\n".join(lines)


def build_excel_sheet_output_path(
    src: Path,
    input_root: Path,
    output_dir: Path,
    output_structure: str,
    forced_ext: Optional[str],
    sheet_name: str,
    include_sheet_name_in_filename: bool,
) -> Path:
    """
    Compute the output Markdown path for a single Excel worksheet export.
    """
    ext = forced_ext if forced_ext not in (None, "") else ".md"
    safe_sheet = sanitize_sheet_name(sheet_name)

    if output_structure == "flat":
        base_name = src.stem
        out_name = f"{base_name}__{safe_sheet}{ext}" if include_sheet_name_in_filename else f"{safe_sheet}{ext}"
        return output_dir / out_name

    rel = src.relative_to(input_root)
    parent = output_dir / rel.parent / src.stem
    out_name = f"{safe_sheet}{ext}" if include_sheet_name_in_filename else apply_output_extension(src.name, ext)
    return parent / out_name


def process_excel_workbook(src: Path, cfg: AppConfig, input_root: Path):
    """
    Process an Excel workbook and write one Markdown file per worksheet when in per-sheet mode.
    """
    messages = []
    written = 0
    failed = 0

    wb = load_workbook(filename=str(src), data_only=True, read_only=True)
    try:
        for ws in wb.worksheets:
            if not cfg.excel_include_empty_sheets and not worksheet_has_data(ws):
                messages.append(f"SKIP (empty sheet): {src} [{ws.title}]")
                continue

            out_path = build_excel_sheet_output_path(
                src=src,
                input_root=input_root,
                output_dir=cfg.output_dir,
                output_structure=cfg.output_structure,
                forced_ext=cfg.output_extension,
                sheet_name=ws.title,
                include_sheet_name_in_filename=cfg.excel_sheet_name_in_filename,
            )

            if out_path.exists() and not cfg.overwrite:
                messages.append(f"SKIP (exists): {src} [{ws.title}]")
                continue

            try:
                ensure_parent_dir(out_path)
                md = worksheet_to_markdown(src, ws, cfg.excel_max_cell_length)
                out_path.write_text(md, encoding="utf-8")
                messages.append(f"OK: {src} [{ws.title}] -> {out_path}")
                written += 1
            except Exception as e:
                messages.append(f"FAIL: {src} [{ws.title}] ({type(e).__name__}: {e})")
                failed += 1
    finally:
        wb.close()

    return written, failed, messages


def local_name(tag: str) -> str:
    """
    Return an XML tag name without namespace decoration.
    """
    return tag.split("}", 1)[-1] if "}" in tag else tag


def element_children_map(elem: ET.Element) -> Dict[str, List[ET.Element]]:
    """
    Build a dictionary of child XML elements grouped by local tag name.
    """
    result: Dict[str, List[ET.Element]] = {}
    for child in list(elem):
        result.setdefault(local_name(child.tag), []).append(child)
    return result


def element_text(elem: Optional[ET.Element]) -> str:
    """Fixed to only return immediate node text, preventing recursive concatenation."""
    if elem is None or elem.text is None:
        return ""
    return elem.text.strip()


def first_text_anywhere(root: ET.Element, names: List[str]) -> str:
    """
    Search the XML tree for the first matching tag name and return its text content.
    """
    wanted = {n.lower() for n in names}
    for elem in root.iter():
        if local_name(elem.tag).lower() in wanted:
            txt = element_text(elem)
            if txt:
                return txt
    return ""


def find_execution_nodes(root: ET.Element, names: List[str]) -> List[ET.Element]:
    """
    Find XML elements whose local tag names match the supplied list of candidate names.
    """
    wanted = {n.lower() for n in names}
    found = []
    for elem in root.iter():
        if local_name(elem.tag).lower() in wanted:
            found.append(elem)
    return found


def parse_datetime_text(text: str) -> str:
    """
    Attempt to normalize known datetime formats while preserving unrecognized strings.
    """
    if not text:
        return ""
    candidates = [
        "%Y-%m-%dT%H:%M:%S.%f%z",
        "%Y-%m-%dT%H:%M:%S%z",
        "%Y-%m-%d %H:%M:%S",
        "%m/%d/%Y %H:%M:%S",
    ]
    for fmt in candidates:
        try:
            dt = datetime.strptime(text, fmt)
            return dt.strftime("%b %-d %Y %-I:%M%p %z") if "%" in fmt else dt.strftime("%b %-d %Y %-I:%M%p")
        except Exception:
            pass
    return text


def normalize_bool_text(value: str) -> str:
    """
    Normalize boolean-like text values to Yes or No when possible.
    """
    v = (value or "").strip().lower()
    if v in {"true", "1", "yes", "y"}:
        return "Yes"
    if v in {"false", "0", "no", "n", ""}:
        return "No"
    return value



# -----------------------------------------------------------------------------
# TestSavvy XML extraction helpers
# -----------------------------------------------------------------------------
def extract_runtime_variables(root: ET.Element) -> List[Tuple[str, str]]:
    """
    Extract runtime variables from TestSavvy runtime variable nodes and attribute-based variable entries.
    """
    vars_out: List[Tuple[str, str]] = []

    candidate_parents = []
    for elem in root.iter():
        name = local_name(elem.tag).lower()
        if name in {"runtimevariables", "runtime_variables", "variables", "runtimevars", "variablelist"}:
            candidate_parents.append(elem)

    def add_pair(k: str, v: str):
        if k:
            vars_out.append((k, v))

    for parent in candidate_parents:
        for elem in parent.iter():
            lname = local_name(elem.tag).lower()
            if lname in {"variable", "runtimevariable", "var", "entry"}:
                key = ""
                val = ""

                for attr_name in ["name", "key"]:
                    if elem.attrib.get(attr_name):
                        key = elem.attrib.get(attr_name, "")
                        break
                for attr_name in ["value"]:
                    if elem.attrib.get(attr_name):
                        val = elem.attrib.get(attr_name, "")
                        break

                if not key or not val:
                    children = element_children_map(elem)
                    for kn in ["name", "key", "variableName", "varName"]:
                        if key:
                            break
                        for c in children.get(kn, []):
                            key = element_text(c)
                            if key:
                                break
                    for vn in ["value", "variableValue", "varValue"]:
                        if val:
                            break
                        for c in children.get(vn, []):
                            val = element_text(c)
                            if val:
                                break

                if key:
                    add_pair(key, val)

    if not vars_out:
        for elem in root.iter():
            lname = local_name(elem.tag).lower()
            if lname in {"variable", "runtimevariable", "var"}:
                key = elem.attrib.get("name") or elem.attrib.get("key") or ""
                val = elem.attrib.get("value") or ""
                if key:
                    add_pair(key, val)

    seen = set()
    deduped = []
    for item in vars_out:
        if item not in seen:
            seen.add(item)
            deduped.append(item)
    return deduped


def extract_step_records(scenario_elem: ET.Element) -> List[Dict[str, str]]:
    """
    Extract TestSavvy step rows from scenario XML, including key attributes and a raw attribute inventory.
    """
    steps: List[Dict[str, str]] = []
    for elem in scenario_elem.iter():
        lname = local_name(elem.tag).lower()
        if lname != "step":
            continue

        dataset = elem.attrib.get("dataset", "") or ""
        header = elem.attrib.get("dataset_header", "") or ""
        dataset_header = f"{dataset} / {header}" if dataset and header else (dataset or header)

        row = {
            "#": elem.attrib.get("count_id") or elem.attrib.get("sequence") or elem.attrib.get("seq") or elem.attrib.get("id") or "",
            "Step ID": elem.attrib.get("step_id", ""),
            "Action": elem.attrib.get("action_name") or elem.attrib.get("actionName") or elem.attrib.get("name") or "",
            "Type": elem.attrib.get("action_type") or elem.attrib.get("actionType") or elem.attrib.get("type") or "",
            "English Text": elem.attrib.get("english_text") or elem.attrib.get("englishText") or "",
            "Dataset / Header": dataset_header,
            "Test Condition": elem.attrib.get("test_condition") or elem.attrib.get("testCondition") or "",
            "Encrypted": "Yes" if str(elem.attrib.get("value_encrypted", "")).strip().lower() in {"1", "true", "yes", "y"} else "No",
            "Value": elem.attrib.get("value", ""),
            "Application": elem.attrib.get("application_name", ""),
            "Interface Map": elem.attrib.get("interface_map", ""),
            "Logical Name": elem.attrib.get("logical_name", ""),
            "Interface Element Type": elem.attrib.get("interface_element_type", ""),
            "Tool": elem.attrib.get("toolplugin") or elem.attrib.get("tool") or "",
            "Comments": elem.attrib.get("comments", ""),
            "Action Code": elem.attrib.get("action_code", ""),
            "Raw Attributes": "; ".join(f"{k}={v}" for k, v in elem.attrib.items()),
        }
        steps.append(row)

    deduped = []
    seen = set()
    for row in steps:
        key = tuple(row.values())
        if key not in seen:
            seen.add(key)
            deduped.append(row)
    return deduped


def extract_automation_sequence(root: ET.Element) -> Optional[ET.Element]:
    """
    Locate the top-level TestSavvy automation_sequence element in the XML document.
    """
    for elem in root.iter():
        if local_name(elem.tag).lower() == "automation_sequence":
            return elem
    return None


def extract_scenarios(root: ET.Element) -> List[Dict[str, Any]]:
    """
    Extract TestSavvy scenario metadata and associated step records from the automation sequence.
    """
    scenarios: List[Dict[str, Any]] = []

    auto_seq = extract_automation_sequence(root)
    candidate_root = auto_seq if auto_seq is not None else root

    for elem in candidate_root.iter():
        lname = local_name(elem.tag).lower()
        if lname != "scenario":
            continue

        title = elem.attrib.get("name") or elem.attrib.get("scenarioName") or elem.attrib.get("title") or "Scenario"
        sid = elem.attrib.get("id") or elem.attrib.get("scenarioId") or ""
        path = elem.attrib.get("path") or ""
        auto_scenario_id = elem.attrib.get("automation_sequence_scenarios_id", "")
        comment = elem.attrib.get("comment_field", "")

        steps = extract_step_records(elem)
        scenarios.append({
            "title": title,
            "scenario_id": sid,
            "automation_sequence_scenarios_id": auto_scenario_id,
            "path": path,
            "comment_field": comment,
            "attributes": dict(elem.attrib),
            "steps": steps,
        })

    deduped = []
    seen = set()
    for s in scenarios:
        key = (s["title"], s["scenario_id"], s["path"], len(s["steps"]))
        if key not in seen:
            seen.add(key)
            deduped.append(s)
    return deduped


def detect_source_guide(root: ET.Element, src: Path) -> str:
    """
    Determine the best human-readable source guide or script name for the XML document.
    """
    auto_seq = extract_automation_sequence(root)
    if auto_seq is not None:
        name = auto_seq.attrib.get("name", "").strip()
        if name:
            return name

    for names in [
        ["sourceGuide", "source_guide", "guideName", "guide"],
        ["automationSequenceName", "sequenceName", "scriptName"],
        ["name", "title"],
    ]:
        val = first_text_anywhere(root, names)
        if val:
            return val
    return src.stem


def extract_execution_summary(root: ET.Element, src: Path, scenario_count: int) -> Dict[str, str]:
    """
    Extract execution-level metadata, including Test Case ID from automation_sequence id.
    """
    auto_seq = extract_automation_sequence(root)

    if auto_seq is not None:
        summary = {
            "Source guide": auto_seq.attrib.get("name", "") or detect_source_guide(root, src),
            "Test Case ID": auto_seq.attrib.get("id", ""),
            "Automation Sequence ID": auto_seq.attrib.get("id", ""),
            "Automation Sequence Name": auto_seq.attrib.get("name", ""),
            "Description": auto_seq.attrib.get("description", ""),
            "Run ID": auto_seq.attrib.get("run_id", "") or src.stem,
            "Status": auto_seq.attrib.get("status_name", "") or auto_seq.attrib.get("status", ""),
            "Status Code": auto_seq.attrib.get("status", ""),
            "Execution Time": auto_seq.attrib.get("execution_time", ""),
            "Execution Sent Date Time": auto_seq.attrib.get("executionsentdatetime", ""),
            "Execution Start Time": auto_seq.attrib.get("executionstarttime", ""),
            "Execution End Time": auto_seq.attrib.get("executionendtime", ""),
            "Machine": auto_seq.attrib.get("machine", ""),
            "Machine ID": auto_seq.attrib.get("machine_id", ""),
            "Executed By": auto_seq.attrib.get("executed_by", ""),
            "Executed By ID": auto_seq.attrib.get("executed_by_id", ""),
            "Last Modified By ID": auto_seq.attrib.get("last_modified_by_id", ""),
            "Client ID": auto_seq.attrib.get("client_id", ""),
            "Iterations": auto_seq.attrib.get("iterations", ""),
            "Verdict": auto_seq.attrib.get("verdict", ""),
            "Reason": auto_seq.attrib.get("reason", ""),
            "Scenario Count": str(scenario_count),
        }
        for key in ["Execution Time", "Execution Sent Date Time", "Execution Start Time", "Execution End Time"]:
            if summary.get(key):
                summary[key] = parse_datetime_text(summary[key])
        return summary

    summary = {
        "Source guide": detect_source_guide(root, src),
        "Test Case ID": "",
        "Automation Sequence ID": "",
        "Automation Sequence Name": "",
        "Description": "",
        "Run ID": first_text_anywhere(root, ["runId", "run_id", "executionId", "execution_id", "id"]) or src.stem,
        "Status": first_text_anywhere(root, ["status", "executionStatus", "runStatus"]),
        "Status Code": "",
        "Execution Time": first_text_anywhere(root, ["executionTime", "runTime", "startTime", "executionStartTime"]),
        "Execution Sent Date Time": "",
        "Execution Start Time": "",
        "Execution End Time": "",
        "Machine": first_text_anywhere(root, ["machine", "machineName", "host", "computerName"]),
        "Machine ID": "",
        "Executed By": first_text_anywhere(root, ["executedBy", "user", "username", "runBy"]),
        "Executed By ID": "",
        "Last Modified By ID": "",
        "Client ID": "",
        "Iterations": "",
        "Verdict": "",
        "Reason": "",
        "Scenario Count": str(scenario_count),
    }
    if summary["Execution Time"]:
        summary["Execution Time"] = parse_datetime_text(summary["Execution Time"])
    return summary



# -----------------------------------------------------------------------------
# Full XML inventory support for completeness / NotebookLM ingestion
# -----------------------------------------------------------------------------
def xml_inventory_rows(root: ET.Element) -> List[Dict[str, str]]:
    """
    Flatten the entire XML tree into rows containing node path, tag name, depth, attributes, and text.
    """
    rows: List[Dict[str, str]] = []

    def visit(elem: ET.Element, path: str, depth: int):
        tag_name = local_name(elem.tag)
        current_path = f"{path}/{tag_name}" if path else tag_name
        # Changed text to only pull immediate node content to avoid unreadable concatenated strings
        text = (elem.text or "").strip()
        rows.append({
            "Path": current_path,
            "Node": tag_name,
            "Depth": str(depth),
            "Attributes": "; ".join(f"{k}={v}" for k, v in elem.attrib.items()),
            "Text": text[:4000],
        })
        for child in list(elem):
            visit(child, current_path, depth + 1)

    visit(root, "", 0)
    return rows


def xml_inventory_markdown(root: ET.Element) -> List[str]:
    """
    Render the flattened XML inventory rows as a Markdown table.
    """
    rows = xml_inventory_rows(root)
    lines = [
        "XML Inventory",
        "| Path | Node | Depth | Attributes | Text |",
        "| --- | --- | --- | --- | --- |",
    ]
    for row in rows:
        lines.append(
            f"| {markdown_escape(row['Path'], 3000)} | "
            f"{markdown_escape(row['Node'], 500)} | "
            f"{markdown_escape(row['Depth'], 50)} | "
            f"{markdown_escape(row['Attributes'], 8000)} | "
            f"{markdown_escape(row['Text'], 4000)} |"
        )
    lines.append("")
    return lines


def scenario_attributes_table_markdown(attributes: Dict[str, str]) -> List[str]:
    """
    Render scenario or automation-sequence attributes as a two-column Markdown table.
    """
    lines = [
        "| Attribute | Value |",
        "| --- | --- |",
    ]
    for key, value in attributes.items():
        lines.append(f"| {markdown_escape(key, 500)} | {markdown_escape(value, 3000)} |")
    lines.append("")
    return lines


def steps_table_markdown(steps: List[Dict[str, str]]) -> List[str]:
    """
    Render extracted TestSavvy step records as a detailed Markdown table.
    """
    header = [
        "#", "Step ID", "Action", "Type", "English Text", "Dataset / Header",
        "Test Condition", "Encrypted", "Value", "Application", "Interface Map",
        "Logical Name", "Interface Element Type", "Tool", "Comments", "Action Code",
        "Raw Attributes"
    ]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * len(header)) + " |",
    ]
    for step in steps:
        row = [markdown_escape(step.get(col, ""), 8000) for col in header]
        lines.append("| " + " | ".join(row) + " |")
    if not steps:
        lines.append("| " + " | ".join([""] * len(header)) + " |")
    return lines



# -----------------------------------------------------------------------------
# Enhanced TestSavvy Markdown report generation
# -----------------------------------------------------------------------------
def testsavvy_execution_report_markdown(src: Path, root: ET.Element, include_xml_inventory: bool = True) -> str:
    """
    Build the enhanced NotebookLM-ready TestSavvy Markdown report from parsed XML.
    """
    scenarios = extract_scenarios(root)
    runtime_vars = extract_runtime_variables(root)
    summary = extract_execution_summary(root, src, len(scenarios))
    auto_seq = extract_automation_sequence(root)

    lines: List[str] = []
    lines.append(f"# {markdown_escape(src.name, 500)}")
    lines.append("")
    lines.append("Source guide")
    lines.append(f"{markdown_escape(summary['Source guide'], 1000)}")
    lines.append("")

    lines.append("Execution Summary")
    if summary["Test Case ID"]:
        lines.append(f"Test Case ID: {markdown_escape(summary['Test Case ID'], 1000)}")
    if summary["Automation Sequence ID"]:
        lines.append(f"Automation Sequence ID: {markdown_escape(summary['Automation Sequence ID'], 1000)}")
    if summary["Automation Sequence Name"]:
        lines.append(f"Automation Sequence Name: {markdown_escape(summary['Automation Sequence Name'], 1000)}")
    if summary["Description"]:
        lines.append(f"Description: {markdown_escape(summary['Description'], 1000)}")
    lines.append(f"Run ID: {markdown_escape(summary['Run ID'], 1000)}")
    if summary["Status"]:
        lines.append(f"Status: {markdown_escape(summary['Status'], 1000)}")
    if summary["Status Code"]:
        lines.append(f"Status Code: {markdown_escape(summary['Status Code'], 1000)}")
    if summary["Execution Time"]:
        lines.append(f"Execution Time: {markdown_escape(summary['Execution Time'], 1000)}")
    if summary["Execution Sent Date Time"]:
        lines.append(f"Execution Sent Date Time: {markdown_escape(summary['Execution Sent Date Time'], 1000)}")
    if summary["Execution Start Time"]:
        lines.append(f"Execution Start Time: {markdown_escape(summary['Execution Start Time'], 1000)}")
    if summary["Execution End Time"]:
        lines.append(f"Execution End Time: {markdown_escape(summary['Execution End Time'], 1000)}")
    if summary["Machine"]:
        lines.append(f"Machine: {markdown_escape(summary['Machine'], 1000)}")
    if summary["Machine ID"]:
        lines.append(f"Machine ID: {markdown_escape(summary['Machine ID'], 1000)}")
    if summary["Executed By"]:
        lines.append(f"Executed By: {markdown_escape(summary['Executed By'], 1000)}")
    if summary["Executed By ID"]:
        lines.append(f"Executed By ID: {markdown_escape(summary['Executed By ID'], 1000)}")
    if summary["Last Modified By ID"]:
        lines.append(f"Last Modified By ID: {markdown_escape(summary['Last Modified By ID'], 1000)}")
    if summary["Client ID"]:
        lines.append(f"Client ID: {markdown_escape(summary['Client ID'], 1000)}")
    if summary["Iterations"]:
        lines.append(f"Iterations: {markdown_escape(summary['Iterations'], 1000)}")
    if summary["Verdict"]:
        lines.append(f"Verdict: {markdown_escape(summary['Verdict'], 1000)}")
    if summary["Reason"]:
        lines.append(f"Reason: {markdown_escape(summary['Reason'], 1000)}")
    lines.append(f"Scenario Count: {markdown_escape(summary['Scenario Count'], 1000)}")
    lines.append("")

    if auto_seq is not None:
        lines.append("Automation Sequence Attributes")
        lines.extend(scenario_attributes_table_markdown(dict(auto_seq.attrib)))

        machine_settings = None
        for child in list(auto_seq):
            if local_name(child.tag).lower() == "machine_settings":
                machine_settings = child
                break
        if machine_settings is not None:
            lines.append("Machine Settings Attributes")
            lines.extend(scenario_attributes_table_markdown(dict(machine_settings.attrib)))

    lines.append("Runtime Variables")
    if runtime_vars:
        for key, val in runtime_vars:
            lines.append(f"{markdown_escape(key, 500)}: {markdown_escape(val, 1000)}")
    else:
        lines.append("None")
    lines.append("")

    lines.append("Scenarios and Steps")
    if not scenarios:
        lines.append("No scenarios were detected.")
        lines.append("")
    else:
        for idx, scenario in enumerate(scenarios, start=1):
            lines.append(f"{idx}. {markdown_escape(scenario['title'], 1000)}")
            if scenario.get("scenario_id"):
                lines.append(f"Scenario ID: {markdown_escape(scenario['scenario_id'], 1000)}")
            if scenario.get("automation_sequence_scenarios_id"):
                lines.append(f"Automation Sequence Scenario ID: {markdown_escape(scenario['automation_sequence_scenarios_id'], 1000)}")
            if scenario.get("path"):
                lines.append(f"Path: {markdown_escape(scenario['path'], 1000)}")
            if scenario.get("comment_field"):
                lines.append(f"Comment: {markdown_escape(scenario['comment_field'], 1000)}")
            lines.append("Scenario Attributes")
            lines.extend(scenario_attributes_table_markdown(scenario.get("attributes", {})))
            lines.extend(steps_table_markdown(scenario.get("steps", [])))
            lines.append("")

    if include_xml_inventory:
        lines.extend(xml_inventory_markdown(root))

    return "\n".join(lines)


def xml_to_markdown(src: Path, cfg: AppConfig) -> str:
    """
    Dispatch XML conversion to either raw Markdown mode or enhanced TestSavvy report mode.
    """
    tree = ET.parse(str(src), parser=ET.XMLParser())
    root = tree.getroot()
    if cfg.xml_output_mode == "raw_markdown":
        xml_text = pretty_print_xml(src, cfg.indent, cfg.xml_encoding, cfg.preserve_xml_declaration).decode(cfg.xml_encoding, errors="replace")
        return wrap_xml_markdown(src, xml_text)
    return testsavvy_execution_report_markdown(src, root, cfg.include_xml_inventory)


def process_standard_file(src: Path, out_path: Path, cfg: AppConfig):
    """
    Process JSON or XML source files and write the configured output artifact.
    """
    if out_path.exists() and not cfg.overwrite:
        return True, f"SKIP (exists): {src}"

    try:
        ensure_parent_dir(out_path)
        ext = src.suffix.lower().lstrip(".")

        if ext == "json":
            pretty = pretty_print_json(src, cfg.indent, cfg.json_sort_keys)
            if (cfg.output_extension or "").lower() == ".md":
                pretty = wrap_json_markdown(src, pretty)
            out_path.write_text(pretty, encoding="utf-8")

        elif ext == "xml":
            if (cfg.output_extension or "").lower() == ".md":
                md = xml_to_markdown(src, cfg)
                out_path.write_text(md, encoding="utf-8")
            else:
                out_path.write_bytes(pretty_print_xml(src, cfg.indent, cfg.xml_encoding, cfg.preserve_xml_declaration))

        else:
            return True, f"SKIP (unsupported ext): {src}"

        return True, f"OK: {src} -> {out_path}"
    except Exception as e:
        return False, f"FAIL: {src} ({type(e).__name__}: {e})"


def build_combined_section_for_file(src: Path, cfg: AppConfig) -> str:
    """
    Build one Markdown section for an XML or JSON file when combined output is enabled.
    """
    ext = src.suffix.lower().lstrip(".")
    if ext == "xml":
        body = xml_to_markdown(src, cfg).strip()
        return body
    if ext in {"json"}:
        body = wrap_json_markdown(src, pretty_print_json(src, cfg.indent, cfg.json_sort_keys)).strip()
        return body
    raise ValueError(f"Unsupported combined mode source: {src}")


def build_excel_combined_section(src: Path, sheet_name: str, content: str) -> str:
    """
    Build one Markdown section for an Excel worksheet when combined output is enabled.
    """
    return f"# {src.name} - {sheet_name}\n\n{content.strip()}\n"



# -----------------------------------------------------------------------------
# Command-line entry point
# -----------------------------------------------------------------------------
def main() -> int:
    """
    Parse command-line arguments, load configuration, process all source files, and print a summary.
    """
    ap = argparse.ArgumentParser(description="Process TestSavvy XML to Markdown reports and Excel tabs to Markdown.")
    ap.add_argument("--config", required=True, help="Path to config JSON")
    args = ap.parse_args()

    cfg = load_config(Path(args.config))
    cfg.output_dir.mkdir(parents=True, exist_ok=True)

    total_considered = 0
    written = 0
    failed = 0
    combined_chunks: List[str] = []

    for input_dir in cfg.input_dirs:
        if not input_dir.exists():
            print(f"WARN: input dir does not exist: {input_dir}")
            continue

        for src in iter_files(input_dir, cfg.recursive, cfg.file_types):
            if is_excluded(src, cfg.exclude_patterns):
                print(f"SKIP (excluded): {src}")
                continue

            total_considered += 1
            ext = src.suffix.lower().lstrip(".")

            if ext in {"xlsx", "xlsm"}:
                wb = load_workbook(filename=str(src), data_only=True, read_only=True)
                try:
                    for ws in wb.worksheets:
                        if not cfg.excel_include_empty_sheets and not worksheet_has_data(ws):
                            print(f"SKIP (empty sheet): {src} [{ws.title}]")
                            continue

                        md = worksheet_to_markdown(src, ws, cfg.excel_max_cell_length)

                        if cfg.xml_markdown_mode == "combined":
                            combined_chunks.append(build_excel_combined_section(src, ws.title, md))
                            print(f"OK: {src} [{ws.title}] -> [combined markdown buffer]")
                            written += 1
                            continue

                        out_path = build_excel_sheet_output_path(
                            src=src,
                            input_root=input_dir,
                            output_dir=cfg.output_dir,
                            output_structure=cfg.output_structure,
                            forced_ext=cfg.output_extension,
                            sheet_name=ws.title,
                            include_sheet_name_in_filename=cfg.excel_sheet_name_in_filename,
                        )
                        if out_path.exists() and not cfg.overwrite:
                            print(f"SKIP (exists): {src} [{ws.title}]")
                            continue
                        try:
                            ensure_parent_dir(out_path)
                            out_path.write_text(md, encoding="utf-8")
                            print(f"OK: {src} [{ws.title}] -> {out_path}")
                            written += 1
                        except Exception as e:
                            print(f"FAIL: {src} [{ws.title}] ({type(e).__name__}: {e})")
                            failed += 1
                finally:
                    wb.close()
                continue

            if cfg.xml_markdown_mode == "combined" and ext in {"xml", "json"} and (cfg.output_extension or "").lower() == ".md":
                try:
                    combined_chunks.append(build_combined_section_for_file(src, cfg))
                    print(f"OK: {src} -> [combined markdown buffer]")
                    written += 1
                except Exception as e:
                    print(f"FAIL: {src} ({type(e).__name__}: {e})")
                    failed += 1
                continue

            out_path = compute_output_path(src, input_dir, cfg.output_dir, cfg.output_structure, cfg.output_extension)
            success, msg = process_standard_file(src, out_path, cfg)
            print(msg)
            if success and msg.startswith("OK:"):
                written += 1
            if not success:
                failed += 1

    if combined_chunks:
        combined_path = cfg.output_dir / cfg.combined_output_filename
        try:
            ensure_parent_dir(combined_path)
            combined_text = "\n\n---\n\n".join(chunk.strip() for chunk in combined_chunks if chunk.strip()) + "\n"
            combined_path.write_text(combined_text, encoding="utf-8")
            print(f"OK: wrote combined markdown -> {combined_path}")
        except Exception as e:
            print(f"FAIL: combined markdown ({type(e).__name__}: {e})")
            failed += 1

    print("\n=== Summary ===")
    print(f"Total considered: {total_considered}")
    print(f"Outputs written:  {written}")
    print(f"Failed:          {failed}")

    return 0 if failed == 0 else 2


if __name__ == "__main__":
    main()
